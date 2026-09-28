import openpyxl

from ais_extractor import build_ais_workbook, extract_tables, read_ais_json


def test_extracts_nested_reported_data_into_workbook():
    payload = read_ais_json(
        b'{"ais":{"tds":[{"amount":100,"deductor":{"name":"Acme"}}],"sft":[{"code":"S1"}]}}'
    )
    tables = extract_tables(payload)
    assert tables["ais.tds"][0]["deductor.name"] == "Acme"
    output, counts = build_ais_workbook(payload)
    assert counts == {"ais.tds": 1, "ais.sft": 1}
    workbook = openpyxl.load_workbook(__import__("io").BytesIO(output), data_only=True)
    assert workbook["Index"].max_row == 3
    assert workbook["tds"].cell(2, 1).value == 100


def test_rejects_invalid_json():
    try:
        read_ais_json(b"not json")
    except ValueError as error:
        assert "Invalid JSON" in str(error)
    else:
        raise AssertionError("Expected invalid JSON to be rejected")
