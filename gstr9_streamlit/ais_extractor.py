"""Convert a readable AIS JSON download into a reviewable Excel workbook.

The Income Tax portal's AIS JSON schema changes over time.  Rather than rely on
hard-coded section names, this module finds every list of JSON objects and
exports each list as a separate worksheet.  This retains the reported data
even when the portal adds a new AIS category.
"""

from __future__ import annotations

import json
import re
from collections import defaultdict
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


def read_ais_json(file_bytes: bytes) -> Any:
    """Parse a UTF-8/UTF-8 BOM JSON file and return its decoded contents."""
    try:
        return json.loads(file_bytes.decode("utf-8-sig"))
    except UnicodeDecodeError as exc:
        raise ValueError("This file is not a readable UTF-8 JSON file.") from exc
    except json.JSONDecodeError as exc:
        raise ValueError(
            f"Invalid JSON (line {exc.lineno}, column {exc.colno})."
        ) from exc


def _scalar(value: Any) -> Any:
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return json.dumps(value, ensure_ascii=False, separators=(",", ":"))


def _flatten_record(record: dict[str, Any], prefix: str = "") -> dict[str, Any]:
    flattened: dict[str, Any] = {}
    for key, value in record.items():
        column = f"{prefix}.{key}" if prefix else str(key)
        if isinstance(value, dict):
            flattened.update(_flatten_record(value, column))
        else:
            flattened[column] = _scalar(value)
    return flattened


def extract_tables(payload: Any) -> dict[str, list[dict[str, Any]]]:
    """Return object-list tables keyed by their JSON path.

    Lists consisting purely of objects are the normal representation of AIS
    reported-data rows. Nested object lists are collected too, so no category
    is silently omitted.
    """
    tables: dict[str, list[dict[str, Any]]] = defaultdict(list)

    def walk(value: Any, path: str) -> None:
        if isinstance(value, list):
            objects = [item for item in value if isinstance(item, dict)]
            if objects:
                tables[path or "Reported Data"].extend(
                    _flatten_record(item) for item in objects
                )
            for item in value:
                if isinstance(item, (dict, list)):
                    walk(item, path)
        elif isinstance(value, dict):
            for key, child in value.items():
                child_path = f"{path}.{key}" if path else str(key)
                if isinstance(child, (dict, list)):
                    walk(child, child_path)

    walk(payload, "")
    return dict(tables)


def _sheet_name(path: str, used: set[str]) -> str:
    name = re.sub(r"[\\/*?:\[\]]", "_", path.split(".")[-1]).strip() or "Reported Data"
    name = name[:31]
    candidate, suffix = name, 2
    while candidate.lower() in used:
        ending = f"_{suffix}"
        candidate = f"{name[:31-len(ending)]}{ending}"
        suffix += 1
    used.add(candidate.lower())
    return candidate


def build_ais_workbook(payload: Any) -> tuple[bytes, dict[str, int]]:
    """Create an XLSX workbook with an index and one sheet per data category."""
    tables = extract_tables(payload)
    if not tables:
        raise ValueError("No reported-data lists were found in this JSON file.")

    workbook = Workbook()
    index = workbook.active
    index.title = "Index"
    index.append(["AIS JSON path", "Worksheet", "Reported rows"])
    header_fill = PatternFill("solid", fgColor="1A56DB")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in index[1]:
        cell.fill, cell.font = header_fill, header_font

    used = {"index"}
    row_counts: dict[str, int] = {}
    for path, rows in tables.items():
        sheet_name = _sheet_name(path, used)
        sheet = workbook.create_sheet(sheet_name)
        columns = list(dict.fromkeys(key for row in rows for key in row))
        sheet.append(columns)
        for cell in sheet[1]:
            cell.fill, cell.font = header_fill, header_font
        for row in rows:
            sheet.append([row.get(column) for column in columns])
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for column_number, column in enumerate(columns, 1):
            longest = max(
                [len(str(column))] + [len(str(row.get(column, ""))) for row in rows]
            )
            sheet.column_dimensions[get_column_letter(column_number)].width = min(
                max(longest + 2, 12), 45
            )
        index.append([path, sheet_name, len(rows)])
        row_counts[path] = len(rows)

    index.freeze_panes = "A2"
    index.auto_filter.ref = index.dimensions
    for column, width in {"A": 55, "B": 28, "C": 16}.items():
        index.column_dimensions[column].width = width
    output = BytesIO()
    workbook.save(output)
    return output.getvalue(), row_counts
